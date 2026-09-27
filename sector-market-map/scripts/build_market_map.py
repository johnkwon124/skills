"""Template: build a three-sheet market map workbook.

All companies below are fictional placeholders. Replace COMPANIES and LAYERS
with real research. Requires: pip install openpyxl
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SECTOR = "Example Sector"
UPDATED = "YYYY-MM"

NAVY, BLUE, ACCENT, LIGHT, EDGE, WHITE = "1F2D3D", "2E4A6B", "3B82F6", "F1F5F9", "CBD5E1", "FFFFFF"
GRADE_FILL = {"A": "DCFCE7", "B": "FEF3C7", "C": "F3F4F6", "Pass": "FEE2E2", "EXITED": "D1D5DB"}


def fill(c): return PatternFill("solid", fgColor=c)
def border():
    s = Side(style="thin", color=EDGE)
    return Border(left=s, right=s, top=s, bottom=s)
def wrap(h="left"): return Alignment(horizontal=h, vertical="center", wrap_text=True)


def grade(total, exited=False):
    if exited: return "EXITED"
    if total >= 13: return "A"
    if total >= 9: return "B"
    if total >= 5: return "C"
    return "Pass"


# name, founded, HQ, stage, funding, layer, customer, model, product,
# investors, inv_score, founders, founder_score, partners, partner_score, exited, notes
COMPANIES = [
    ("Acme Grid", 2022, "City A", "Series A", "$30M", "Platform", "Operators", "SaaS",
     "Placeholder description of product and differentiator.",
     "Example Ventures (lead), Sample Capital", 4,
     "Placeholder founder background.", 4,
     "Announced pilot with Example Utility", 3, False, "Placeholder scoring rationale."),
    ("Example Robotics", 2020, "City B", "Seed", "$8M", "Hardware", "Enterprise", "Hardware + service",
     "Placeholder description.", "Demo Fund", 2, "Placeholder background.", 3,
     "None announced", 1, False, "Thin public information, scored low."),
    ("Sample Power Co", 2015, "City C", "Acquired", "N/A", "Infrastructure", "Utilities", "Licensing",
     "Placeholder description.", "N/A", 3, "Placeholder background.", 3,
     "Acquired by Example Corp", 3, True, "Exit kept as a benchmark."),
]

LAYERS = [  # layer, players, description, investment relevance
    ("Platform", "Acme Grid", "Placeholder layer description.", "Placeholder relevance."),
    ("Hardware", "Example Robotics", "Placeholder layer description.", "Placeholder relevance."),
]

OBSERVATIONS = [
    ("1. Market timing", "Placeholder observation."),
    ("2. Open questions for diligence", "Placeholder question."),
]

NEXT_STEP = "Request intro meeting"


def title(ws, cols, text, sub):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    c = ws.cell(1, 1, text)
    c.font, c.fill, c.alignment = Font(name="Arial", bold=True, size=14, color=WHITE), fill(NAVY), wrap("center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    c = ws.cell(2, 1, sub)
    c.font, c.fill, c.alignment = Font(name="Arial", size=9, color=WHITE), fill(BLUE), wrap("center")


def header(ws, row, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row, i, n)
        c.font, c.fill, c.alignment, c.border = Font(name="Arial", bold=True, size=9, color=WHITE), fill(ACCENT), wrap("center"), border()
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# Sheet 1: long list
ws = wb.active
ws.title = "Long List"
ws.sheet_view.showGridLines = False
cols = ["Company", "Founded", "HQ", "Stage", "Total Funding", "Value Chain Position", "Customer Segment",
        "Business Model", "Core Product / Differentiator", "Key Investors", "Investor Score (1-5)",
        "Founder Background", "Founder Score (1-5)", "Key Partnerships", "Partnership Score (1-5)",
        "Total Score (/15)", "Rating", "Evaluation Notes"]
widths = [22, 9, 14, 10, 14, 16, 16, 14, 38, 32, 10, 32, 10, 32, 10, 9, 10, 44]
title(ws, len(cols), f"{SECTOR}: Market Map and Rating",
      f"Updated {UPDATED} | Source: public information only")
header(ws, 4, cols, widths)
ws.freeze_panes = "A5"

rated = []
for r, co in enumerate(COMPANIES, 5):
    total = co[10] + co[12] + co[14]
    g = grade(total, co[15])
    rated.append((co, total, g))
    row = list(co[:15]) + [total, g, co[16]]
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v)
        cell.font, cell.border = Font(name="Arial", size=9, bold=(c in (16, 17))), border()
        cell.alignment = wrap("left" if c in (1, 6, 7, 8, 9, 10, 12, 14, 18) else "center")
        cell.fill = fill(GRADE_FILL[g])
    ws.row_dimensions[r].height = 80

# Sheet 2: competitive landscape
ws2 = wb.create_sheet("Competitive Landscape")
ws2.sheet_view.showGridLines = False
title(ws2, 4, f"{SECTOR}: Competitive Landscape", "Value-chain layers and structural observations")
header(ws2, 3, ["Layer", "Players", "Description", "Investment Relevance"], [28, 24, 46, 46])
for r, row in enumerate(LAYERS, 4):
    for c, v in enumerate(row, 1):
        cell = ws2.cell(r, c, v)
        cell.font, cell.border, cell.alignment = Font(name="Arial", size=9), border(), wrap()
    ws2.row_dimensions[r].height = 60
r0 = len(LAYERS) + 5
ws2.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=4)
c = ws2.cell(r0, 1, "KEY STRUCTURAL OBSERVATIONS")
c.font, c.fill, c.alignment = Font(name="Arial", bold=True, size=10, color=WHITE), fill(BLUE), wrap("center")
for r, (topic, text) in enumerate(OBSERVATIONS, r0 + 1):
    ws2.cell(r, 1, topic).font = Font(name="Arial", bold=True, size=9)
    ws2.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws2.cell(r, 2, text).font = Font(name="Arial", size=9)
    for c in (1, 2):
        ws2.cell(r, c).alignment, ws2.cell(r, c).border = wrap(), border()
    ws2.row_dimensions[r].height = 45

# Sheet 3: A-grade shortlist
ws3 = wb.create_sheet("A-Grade Shortlist")
ws3.sheet_view.showGridLines = False
title(ws3, 6, "A-Grade Shortlist", "Priority outreach targets (total score 13 or higher)")
header(ws3, 3, ["Company", "Score", "Stage / Funding", "Why A-Grade", "Key Risks", "Recommended Next Step"],
       [22, 9, 22, 50, 40, 30])
r = 4
for co, total, g in rated:
    if g == "A":
        for c, v in enumerate([co[0], total, f"{co[3]} / {co[4]}", co[16], "Fill from research", NEXT_STEP], 1):
            cell = ws3.cell(r, c, v)
            cell.font, cell.border, cell.alignment = Font(name="Arial", size=9), border(), wrap()
        ws3.row_dimensions[r].height = 60
        r += 1
b_names = ", ".join(co[0] for co, _, g in rated if g == "B") or "none"
ws3.cell(r + 1, 1, f"Monitoring queue (B-grade): {b_names}").font = Font(name="Arial", italic=True, size=9)

wb.save(f"{SECTOR.replace(' ', '_')}_MarketMap.xlsx")
print("saved")
