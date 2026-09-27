#!/usr/bin/env python3
"""Mechanical lint for the document-fact-gate pipeline. Deterministic, no model.

Usage:
    python3 lint.py <file.docx|file.txt|file.md> [--require-section NAME ...] [--json]

Exit code 0 means no FAIL findings, 1 means at least one FAIL.
WARN findings do not change the exit code but must be resolved in verification.
"""
import sys, re, json, argparse

REL_DATE = re.compile(
    r"\b(this|next|last)\s+(year|summer|spring|fall|autumn|winter|quarter|month|half)\b"
    r"|\blater this year\b|\bearly next year\b|\bby year[- ]?end\b"
    r"|\bin the coming (months|weeks|years)\b|\bin the near (term|future)\b"
    r"|\brecently\b|\bsoon\b",
    re.IGNORECASE)

EM_DASH = re.compile("\u2014")

PROCESS = re.compile(
    r"\boriginally (i )?(wrote|had|said)\b|\bon re-?check(ing)?\b|\bafter re-?verif\w+\b"
    r"|\brevision history\b|\bin this revision\b|\bafter (several|a few|multiple) (passes|rounds|revisions)\b"
    r"|\bwhile reviewing,? (i|we) (found|noticed)\b|\bexcluded because (no|the) source\b",
    re.IGNORECASE)

FIGURE = re.compile(
    r"\$\s?\d|\d+(\.\d+)?\s?%|\b\d{1,3}(,\d{3})+\b"
    r"|\b\d+(\.\d+)?\s?(MW|GW|kW|MWh|GWh|million|billion)\b", re.IGNORECASE)
SOURCE_MARK = re.compile(
    r"\(stated by company\)|\(company-stated\)|\(unverified\)|\(source[:\s]|\[S\d+\]|according to|per\s+[A-Z]",
    re.IGNORECASE)


def extract_text(path):
    if path.lower().endswith(".docx"):
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        for section in doc.sections:
            for p in section.footer.paragraphs:
                parts.append(p.text)
        return "\n".join(parts)
    with open(path, encoding="utf-8") as f:
        return f.read()


def lint(text, required_sections=()):
    findings = []

    def ctx(a, b):
        return re.sub(r"\s+", " ", text[max(0, a - 40):b + 40]).strip()

    for m in REL_DATE.finditer(text):
        findings.append({"rule": "REL_DATE", "severity": "FAIL",
                         "match": m.group(0), "context": ctx(m.start(), m.end())})
    for m in EM_DASH.finditer(text):
        findings.append({"rule": "EM_DASH", "severity": "FAIL",
                         "match": "long dash", "context": ctx(m.start(), m.end())})
    for m in PROCESS.finditer(text):
        findings.append({"rule": "PROCESS_NARRATION", "severity": "WARN",
                         "match": m.group(0), "context": ctx(m.start(), m.end())})

    for line in text.splitlines():
        if FIGURE.search(line) and not SOURCE_MARK.search(line):
            findings.append({"rule": "UNSOURCED_FIGURE", "severity": "WARN",
                             "match": line.strip()[:140]})

    low = text.lower()
    for h in required_sections:
        if h.lower() not in low:
            findings.append({"rule": "MISSING_SECTION", "severity": "FAIL", "match": h})

    return {
        "summary": {
            "fail": sum(1 for f in findings if f["severity"] == "FAIL"),
            "warn": sum(1 for f in findings if f["severity"] == "WARN"),
            "tags": {
                "stated_by_company": len(re.findall(r"\(stated by company\)", text, re.I)),
                "unverified": len(re.findall(r"\(unverified\)", text, re.I)),
            },
        },
        "findings": findings,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--require-section", action="append", default=[])
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    report = lint(extract_text(args.file), args.require_section)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        s = report["summary"]
        print(f"FAIL: {s['fail']}  WARN: {s['warn']}  "
              f"stated-by-company: {s['tags']['stated_by_company']}  unverified: {s['tags']['unverified']}")
        for f in report["findings"]:
            print(f"[{f['severity']}] {f['rule']}: {f.get('context', f['match'])}")
    sys.exit(1 if report["summary"]["fail"] > 0 else 0)


if __name__ == "__main__":
    main()
